# -*- coding: utf-8 -*-
"""Excel (.xlsx) export destination.

Appends normalised transactions to an ``.xlsx`` workbook (via the optional
``openpyxl`` library).  Like the CSV destination, appending is naturally
idempotent because the engine only ever hands over not-yet-delivered rows.
"""

from __future__ import annotations

import os
from typing import List

from bankhub.errors import MissingDependencyError
from bankhub.models import PushResult, Transaction
from bankhub.registry import register_destination
from bankhub.destinations.base import Destination

_FIELDS = ["external_id", "date", "amount", "currency", "payee", "notes",
           "category", "account", "source", "status"]


def _row(txn: Transaction) -> list:
    return [txn.external_id, txn.date.isoformat(), float(txn.amount), txn.currency,
            txn.payee, txn.notes, txn.category or "", txn.target_account,
            txn.source, txn.status]


@register_destination("xlsx")
class XlsxDestination(Destination):
    """Append transactions to an Excel workbook.

    Options
    -------
    file
        Output path (default ``export.xlsx``).
    sheet
        Worksheet title (default ``Transactions``).
    mode
        ``"append"`` (default) or ``"overwrite"``.
    """

    requires = "openpyxl"

    def __init__(self, file: str = "export.xlsx", sheet: str = "Transactions",
                 mode: str = "append", **options):
        super().__init__(file=file, sheet=sheet, mode=mode, **options)
        self.file = file
        self.sheet = sheet
        self.mode = mode

    def push(self, transactions: List[Transaction]) -> List[PushResult]:
        if not transactions:
            return []
        try:
            import openpyxl
        except ImportError:
            raise MissingDependencyError("xlsx", "openpyxl")

        directory = os.path.dirname(os.path.abspath(self.file))
        os.makedirs(directory, exist_ok=True)

        if self.mode == "append" and os.path.exists(self.file):
            wb = openpyxl.load_workbook(self.file)
            ws = wb[self.sheet] if self.sheet in wb.sheetnames else wb.create_sheet(self.sheet)
            if ws.max_row == 1 and all(c.value is None for c in ws[1]):
                ws.append(_FIELDS)
        else:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = self.sheet
            ws.append(_FIELDS)

        results = []
        for txn in transactions:
            ws.append(_row(txn))
            results.append(PushResult(external_id=txn.external_id, status="created"))
        wb.save(self.file)
        return results
